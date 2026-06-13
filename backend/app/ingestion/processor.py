"""Document text extraction + the end-to-end ingestion pipeline.

Pipeline: raw content -> chunk -> embed -> store chunks in pgvector.
Used by the Memory API for documents, URLs and notes.
"""
from __future__ import annotations

import io
from typing import Any

from app.db import queries
from app.ingestion.chunker import chunk_text
from app.ingestion.embedder import embed_texts
from app.utils.logger import get_logger

log = get_logger(__name__)


# ---------------------------------------------------------------------------
# Text extraction by file type
# ---------------------------------------------------------------------------
def extract_text(filename: str, data: bytes, content_type: str | None) -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf") or content_type == "application/pdf":
            return _extract_pdf(data)
        if name.endswith(".docx"):
            return _extract_docx(data)
        if name.endswith(".xlsx"):
            return _extract_xlsx(data)
        if name.endswith(".csv"):
            return data.decode("utf-8", errors="replace")
        # txt / md / fallback
        return data.decode("utf-8", errors="replace")
    except Exception as exc:  # pragma: no cover - defensive
        log.exception("Text extraction failed for %s", filename)
        raise ValueError(f"Could not extract text from {filename}: {exc}") from exc


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)


# ---------------------------------------------------------------------------
# Ingestion pipeline
# ---------------------------------------------------------------------------
async def ingest_item(item: dict[str, Any], content: str) -> int:
    """Chunk + embed + store `content` for an existing knowledge_item.

    Returns the number of chunks stored. Updates the item's status.
    """
    item_id = item["id"]
    user_id = item["user_id"]
    try:
        queries.update_knowledge_item(item_id, {"status": "processing"})

        chunks = chunk_text(content)
        if not chunks:
            queries.update_knowledge_item(
                item_id, {"status": "ready", "chunk_count": 0}
            )
            return 0

        # Embed in batches to respect provider limits.
        rows: list[dict[str, Any]] = []
        batch_size = 96
        for start in range(0, len(chunks), batch_size):
            batch = chunks[start : start + batch_size]
            vectors = await embed_texts(batch)
            for offset, (text, vec) in enumerate(zip(batch, vectors)):
                rows.append(
                    {
                        "item_id": item_id,
                        "user_id": user_id,
                        "content": text,
                        "embedding": vec,
                        "chunk_index": start + offset,
                        "metadata": {
                            "title": item.get("title"),
                            "type": item.get("type"),
                            "source_url": item.get("source_url"),
                        },
                    }
                )

        queries.insert_chunks(rows)
        queries.update_knowledge_item(
            item_id, {"status": "ready", "chunk_count": len(rows)}
        )
        return len(rows)
    except Exception as exc:
        log.exception("Ingestion failed for item %s", item_id)
        queries.update_knowledge_item(
            item_id, {"status": "error", "error": str(exc)[:500]}
        )
        raise
